import os

def get_autor_arquivo(item, loc):
    caminho = item.get("dir_atual") or item.get("dir_anterior")
    if not caminho or not os.path.exists(caminho):
        if "autor" in item and item["autor"]:
            return loc.traduzir_metadados(item["autor"], "autor")
        return ""

    autor = ""
    base, ext = os.path.splitext(caminho)
    ext = ext.lower()

    try:
        if ext in [".docx", ".dotx", ".docm", ".dotm"]:
            from docx import Document
            doc = Document(caminho)
            props = doc.core_properties
            autor = props.author or ""

        elif ext == ".doc":
            import olefile
            try:
                with olefile.OleFileIO(caminho) as ole:
                    if ole.exists('\x05SummaryInformation'):
                        props = ole.getproperties('\x05SummaryInformation')
                        autor = props.get(4, loc.get_text("unknown_author"))

                        if isinstance(autor, bytes):
                            autor = autor.decode("latin-1", errors="ignore")

            except Exception as e:
                print(f"Erro ao obter autor de {ext} usando olefile: {e}")
                autor = loc.get_text("unknown_author")

        elif ext in [".xlsx", ".xlsm", ".xltx", ".xltm"]:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(caminho, read_only=True, data_only=True)

                if wb.properties.creator:
                    autor = wb.properties.creator

                wb.close()

            except Exception as xlsx_err:
                print(f"Erro ao ler XLSX com openpyxl: {xlsx_err}")
                try:
                    import olefile
                    with olefile.OleFileIO(caminho) as ole:
                        if ole.exists('\x05SummaryInformation'):
                            props = ole.getproperties('\x05SummaryInformation')
                            autor = props.get(4, loc.get_text("unknown_author"))

                            if isinstance(autor, bytes):
                                autor = autor.decode("latin-1", errors="ignore")

                except Exception as ole_err:
                    print(f"Erro ao ler XLSX com olefile: {ole_err}")
                    autor = loc.get_text("unknown_author")

        elif ext == ".xls":
            import olefile
            try:
                with olefile.OleFileIO(caminho) as ole:
                    if ole.exists('\x05SummaryInformation'):
                        props = ole.getproperties('\x05SummaryInformation')
                        autor = props.get(4, loc.get_text("unknown_author"))

                        if isinstance(autor, bytes):
                            autor = autor.decode("latin-1", errors="ignore")

            except Exception as e:
                print(f"Erro ao obter autor de {ext} usando olefile: {e}")
                autor = loc.get_text("unknown_author")

        elif ext in [".pptx", ".potx", ".ppsx"]:
            from pptx import Presentation
            pres = Presentation(caminho)
            autor = pres.core_properties.author or ""

        elif ext == ".ppt":
            import olefile
            try:
                with olefile.OleFileIO(caminho) as ole:
                    if ole.exists('\x05SummaryInformation'):
                        props = ole.getproperties('\x05SummaryInformation')
                        autor = props.get(4, loc.get_text("unknown_author"))

                        if isinstance(autor, bytes):
                            autor = autor.decode("latin-1", errors="ignore")

            except Exception as e:
                print(f"Erro ao obter autor de {ext} usando olefile: {e}")
                autor = loc.get_text("unknown_author")

        elif ext in [".mdb", ".accdb"]:
            autor = loc.get_text("access")

        elif ext == ".msg":
            import olefile
            try:
                with olefile.OleFileIO(caminho) as ole:
                    if ole.exists('\x05SummaryInformation'):
                        props = ole.getproperties('\x05SummaryInformation')
                        autor = props.get(4, loc.get_text("unknown_author"))

                        if isinstance(autor, bytes):
                            autor = autor.decode("latin-1", errors="ignore")

                    elif ole.exists('__properties_version1.0'):
                        props = ole.getproperties('__properties_version1.0')
                        for prop_id in [0x0C1A, 0x0E04, 0x0042]:
                            if prop_id in props:
                                autor = props[prop_id]
                                if isinstance(autor, bytes):
                                    autor = autor.decode("latin-1", errors="ignore")

                                break

                    else:
                        autor = loc.get_text("outlook_message")

            except Exception as e:
                print(f"Erro ao extrair informações do MSG: {e}")
                autor = loc.get_text("outlook_message")

        elif ext in [".pst", ".ost"]:
            autor = loc.get_text("outlook")

        elif ext == ".pub":
            autor = loc.get_text("publisher")

        elif ext in [".vsd", ".vsdx"]:
            autor = loc.get_text("visio")

        elif ext in [".mpp", ".mpt"]:
            autor = loc.get_text("project")

        elif ext == ".pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(caminho)
            info = reader.metadata
            autor = (info.author or "") if info else ""

        elif ext in [".txt", ".htm", ".html", ".mht", ".mhtml"]:
            autor = loc.get_text("text_html")

        elif ext in [".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".cab"]:
            autor = loc.get_text("compressed_file")

        else:
            autor = loc.get_text("unknown_author")

    except Exception as e:
        print(f"Erro ao obter autor do arquivo {caminho}: {e}")
        pass

    return loc.traduzir_metadados(autor, "autor")

    return autor
